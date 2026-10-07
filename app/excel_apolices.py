"""Atualização segura do Excel operacional preservando colunas manuais."""
from pathlib import Path
from datetime import datetime
import shutil,time,os
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
from .models import PolicyData
ROBOT_MANAGED_COLUMNS=["ÓRGÃO","EMPRESA","Nº CONCORRÊNCIA/EDITAL","PROCESSO SEI","OBJETO","VIGÊNCIA DATA INICIAL","VIGÊNCIA DATA FINAL","VALOR DO PRÊMIO","Nº REGISTRO DA SUSEP","Nº DA LINHA DIGITÁVEL DO BOLETO"]
def backup_file(path:Path,backup_dir:Path)->Path|None:
    if not path.exists():return None
    backup_dir.mkdir(parents=True,exist_ok=True);dest=backup_dir/f"controle_apolices_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx";shutil.copy2(path,dest);return dest

def _key(company,process,number,org):
    c=str(company or "").strip().casefold();p=str(process or "").strip().casefold();n=str(number or "").strip().casefold();o=str(org or "").strip().casefold()
    if c and n and p:return (c,p,n)
    if c and n and o:return (c,n,o)
    return None

def update_workbook(path:Path,data:PolicyData,backup_dir:Path)->str:
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    if path.exists(): backup_file(path,backup_dir)
    if path.exists():wb=load_workbook(path)
    else:
        wb=Workbook();ws=wb.active;ws.title="APÓLICES";ws.append(ROBOT_MANAGED_COLUMNS)
        for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="17365D")
        ws.freeze_panes="A2"
    ws=wb.active
    if ws.max_row<1 or all(ws.cell(1,c).value is None for c in range(1,ws.max_column+1)):raise ValueError("Cabeçalho operacional inválido; workbook não alterado")
    header={str(ws.cell(1,c).value).strip():c for c in range(1,ws.max_column+1) if ws.cell(1,c).value is not None}
    for name in ROBOT_MANAGED_COLUMNS:
        if name not in header:
            col=ws.max_column+1;ws.cell(1,col,name);header[name]=col
    company=str(data.empresa or data.tipo_empresa or "").strip().casefold()
    process=str(data.processo_sei or "").strip().casefold()
    number=str(data.numero_concorrencia_normalizado or "").strip().casefold()
    org=str(data.orgao or "").strip().casefold()
    key=_key(company,process,number,org)
    if not key:raise ValueError("Chave lógica insuficiente; revisão manual necessária")
    matches=[]
    for r in range(2,ws.max_row+1):
        existing_company=str(ws.cell(r,header["EMPRESA"]).value or "").strip().casefold()
        existing_process=str(ws.cell(r,header["PROCESSO SEI"]).value or "").strip().casefold()
        existing_number=str(ws.cell(r,header["Nº CONCORRÊNCIA/EDITAL"]).value or "").strip().casefold()
        existing_org=str(ws.cell(r,header["ÓRGÃO"]).value or "").strip().casefold()
        same_case=bool(company and number and existing_company==company and existing_number in {number,str(data.numero_concorrencia_original or "").strip().casefold()})
        process_match=bool(process and existing_process and process==existing_process)
        fallback_match=bool((not process or not existing_process) and org and existing_org==org)
        if same_case and (process_match or fallback_match):matches.append(r)
    if len(matches)>1:raise ValueError("POSSIVEL_DUPLICATA: múltiplas linhas correspondem à chave; sem alteração")
    row=matches[0] if matches else ws.max_row+1
    values={"ÓRGÃO":data.orgao,"EMPRESA":data.empresa or data.tipo_empresa,"Nº CONCORRÊNCIA/EDITAL":data.numero_concorrencia_original or data.numero_concorrencia_normalizado,"PROCESSO SEI":data.processo_sei,"OBJETO":data.objeto,"VIGÊNCIA DATA INICIAL":data.vigencia_data_inicial,"VIGÊNCIA DATA FINAL":data.vigencia_data_final,"VALOR DO PRÊMIO":data.valor_premio,"Nº REGISTRO DA SUSEP":data.numero_registro_susep,"Nº DA LINHA DIGITÁVEL DO BOLETO":data.linha_digitavel_boleto}
    for name,value in values.items():
        cell=ws.cell(row,header[name])
        if value is not None:cell.value=value
        if name.startswith("VIGÊNCIA"):cell.number_format="dd/mm/yyyy"
        if name=="VALOR DO PRÊMIO":cell.number_format='R$ #,##0.00'
        if name=="Nº DA LINHA DIGITÁVEL DO BOLETO":cell.number_format="@"
    ws.auto_filter.ref=ws.dimensions;ws.freeze_panes=ws.freeze_panes or "A2"
    for name,col in header.items():
        if name in ROBOT_MANAGED_COLUMNS:ws.column_dimensions[ws.cell(1,col).column_letter].width=max(18,min(48,len(name)+4))
    if ws.max_row>=2:
        if ws.tables:
            for table in ws.tables.values():table.ref=ws.dimensions
        else:
            table_headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
            if all(table_headers) and len(set(map(str,table_headers)))==len(table_headers):
                table=Table(displayName="ApolicesTable",ref=ws.dimensions)
                table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
                ws.add_table(table)
    temp=path.with_name(path.stem+".tmp"+path.suffix)
    try:
        wb.save(temp)
        # Validar que o arquivo temporário abre antes da substituição atômica.
        check=load_workbook(temp,read_only=True);check.close()
        os.replace(temp,path)
    except Exception:
        if temp.exists():temp.unlink()
        raise
    return "ATUALIZADO" if matches else "ADICIONADO"
