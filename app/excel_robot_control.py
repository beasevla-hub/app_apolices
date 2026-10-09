"""Planilha técnica de auditoria e idempotência por par documental."""
from pathlib import Path
from datetime import datetime
from uuid import uuid4
import os,json
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
HEADERS=["message_id","uid","group_key","lote","lotes","data_email","remetente","assunto","hash_apolice","hash_boleto","status","tentativas","modelo_ia","data_ultima_tentativa","data_processamento","erro","pasta_destino","id_processamento"]
class RobotControl:
    def __init__(self,path:Path,max_attempts:int=5):self.path=Path(path);self.max_attempts=max_attempts
    def _open(self):
        if self.path.exists():
            wb=load_workbook(self.path)
            if "EMAILS_PROCESSADOS" not in wb.sheetnames:raise ValueError("Aba EMAILS_PROCESSADOS ausente no controle do robô")
            ws=wb["EMAILS_PROCESSADOS"];current={str(ws.cell(1,c).value):c for c in range(1,ws.max_column+1) if ws.cell(1,c).value}
            for name in HEADERS:
                if name not in current:ws.cell(1,ws.max_column+1,name)
            return wb
        wb=Workbook();ws=wb.active;ws.title="EMAILS_PROCESSADOS";ws.append(HEADERS)
        for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="17365D")
        ws.freeze_panes="A2";ws.auto_filter.ref=f"A1:{ws.cell(1,len(HEADERS)).column_letter}1";return wb
    def rows(self):
        wb=self._open();ws=wb["EMAILS_PROCESSADOS"];headers={str(ws.cell(1,c).value):c for c in range(1,ws.max_column+1) if ws.cell(1,c).value};result=[]
        for row in range(2,ws.max_row+1):
            item={key:ws.cell(row,col).value for key,col in headers.items()}
            if any(value is not None for value in item.values()):result.append(item)
        if not self.path.exists():self.path.parent.mkdir(parents=True,exist_ok=True);wb.save(self.path)
        return result
    @staticmethod
    def _group_matches(row,group_key):return (row.get("group_key") or "")== (group_key or "")
    def find(self,message_id:str,uid:str="",group_key:str|None=None):
        return next((row for row in self.rows() if ((message_id and row.get("message_id")==message_id) or (uid and str(row.get("uid"))==str(uid))) and self._group_matches(row,group_key)),None)
    def completed_pair(self,message_id:str,group_key:str,policy_hash:str,bill_hash:str,lot:str|None=None,lotes:list[str]|None=None):
        """Retorna somente SUCESSO com os dois hashes exatos; status/tentativas não são gates."""
        completed=[row for row in self.rows() if row.get("status")=="SUCESSO" and row.get("hash_apolice")==policy_hash and row.get("hash_boleto")==bill_hash]
        def normalize(value):
            text=" ".join(str(value or "").casefold().split())
            return (text[5:].strip() if text.startswith("lote ") else text) or None
        requested=normalize(lot)
        compatible=[row for row in completed if requested is None or normalize(row.get("lote"))==requested]
        if lotes:
            expected={normalize(value) for value in lotes}
            compatible=[]
            for row in completed:
                try:recorded=json.loads(row.get("lotes") or "[]")
                except (TypeError,ValueError):recorded=[]
                if {normalize(value) for value in recorded}==expected:compatible.append(row)
        exact=next((row for row in compatible if row.get("message_id")==message_id and (row.get("group_key") or "")==group_key),None)
        return exact or next(iter(compatible),None)
    def attempts(self,message_id:str,uid:str="",group_key:str|None=None)->int:
        record=self.find(message_id,uid,group_key);return int(record.get("tentativas") or 0) if record else 0
    def begin(self,**data)->str:
        process_id=data.get("id_processamento") or uuid4().hex[:12].upper()
        data.update(status="PROCESSANDO",erro=None,pasta_destino=None,id_processamento=process_id,increment_attempt=True)
        self.record(**data);return process_id
    def record(self,**data):
        self.path.parent.mkdir(parents=True,exist_ok=True);wb=self._open();ws=wb["EMAILS_PROCESSADOS"]
        headers={str(ws.cell(1,col).value):col for col in range(1,ws.max_column+1) if ws.cell(1,col).value}
        mid=data.get("message_id","");uid=data.get("uid","");group_key=data.get("group_key") or ""
        row_num=None
        for idx in range(2,ws.max_row+1):
            same_id=bool(mid and ws.cell(idx,headers["message_id"]).value==mid) or bool(uid and str(ws.cell(idx,headers["uid"]).value)==str(uid))
            existing_group=ws.cell(idx,headers["group_key"]).value or ""
            if same_id and existing_group==group_key:row_num=idx;break
        if row_num is None:row_num=ws.max_row+1
        old_attempt=int(ws.cell(row_num,headers["tentativas"]).value or 0) if row_num<=ws.max_row else 0
        values={key:data.get(key) for key in HEADERS if key in data}
        if data.get("increment_attempt"):values["tentativas"]=old_attempt+1
        elif "tentativas" not in values and old_attempt:values["tentativas"]=old_attempt
        values.setdefault("data_ultima_tentativa",datetime.now().isoformat(timespec="seconds"))
        if values.get("status") in {"PROCESSANDO","SUCESSO"}:values["erro"]=None
        if values.get("status")=="SUCESSO":values["data_processamento"]=datetime.now().isoformat(timespec="seconds")
        for key,value in values.items():
            if key in headers:ws.cell(row_num,headers[key]).value=value
        temp=self.path.with_name(self.path.stem+".tmp"+self.path.suffix)
        try:wb.save(temp);load_workbook(temp,read_only=True).close();os.replace(temp,self.path)
        except Exception:
            if temp.exists():temp.unlink()
            raise
