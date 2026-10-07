from datetime import date
from openpyxl import Workbook,load_workbook
from openpyxl.worksheet.table import Table
import pytest
from app.models import PolicyData
from app.excel_apolices import update_workbook,ROBOT_MANAGED_COLUMNS

def item(process='SEI-1',lot=None):
    return PolicyData(empresa='THI Engenharia',empresa_normalizada='THI',orgao='Sub Parelheiros',orgao_normalizado='SUB PARELHEIROS',numero_concorrencia_original='018/SEME/2026',numero_concorrencia_normalizado='018-SEME-2026',processo_sei=process,lote=lot,vigencia_data_inicial=date(2026,8,1),vigencia_data_final=date(2027,8,1),confianca_geral=.9)
def workbook(path,rows):
    wb=Workbook();ws=wb.active;ws.title='APÓLICES';ws.append(ROBOT_MANAGED_COLUMNS+['STATUS','RESPONSÁVEL'])
    for row in rows:ws.append(row)
    wb.save(path)
def row(process):return ['Sub Parelheiros','THI Engenharia','018/SEME/2026',process,'manual objeto',None,None,None,None,None,None,'PAGO','Bea']
def test_different_sei_does_not_fallback_to_org(tmp_path):
    path=tmp_path/'ops.xlsx';workbook(path,[row('SEI-1')])
    assert update_workbook(path,item('SEI-2'),tmp_path/'backups')=='ADICIONADO'
    ws=load_workbook(path)['APÓLICES'];assert ws.max_row==3 and ws['D2'].value=='SEI-1' and ws['D3'].value=='SEI-2'
def test_same_sei_updates_preserves_manual_and_null(tmp_path):
    path=tmp_path/'ops.xlsx';workbook(path,[row('SEI-1')])
    update_workbook(path,item('SEI-1'),tmp_path/'backups');ws=load_workbook(path)['APÓLICES']
    assert ws.max_row==2 and ws['L2'].value=='PAGO' and ws['M2'].value=='Bea' and ws['E2'].value=='manual objeto'
def test_missing_operational_sheet_fails_without_changing_other_sheet(tmp_path):
    path=tmp_path/'ops.xlsx';wb=Workbook();wb.active.title='Outra';wb.save(path)
    with pytest.raises(ValueError,match='APÓLICES'):update_workbook(path,item(),tmp_path/'backups')
    assert load_workbook(path).sheetnames==['Outra']
def test_missing_sei_uses_org_fallback_and_multiple_matches_fail(tmp_path):
    path=tmp_path/'ops.xlsx';workbook(path,[row(None)])
    update_workbook(path,item(None),tmp_path/'backups');assert load_workbook(path)['APÓLICES'].max_row==2
    path2=tmp_path/'dup.xlsx';workbook(path2,[row(None),row(None)])
    with pytest.raises(ValueError,match='POSSIVEL_DUPLICATA'):update_workbook(path2,item(None),tmp_path/'backups')

def test_existing_table_filter_formula_and_manual_column_are_untouched(tmp_path):
    path=tmp_path/'table.xlsx';workbook(path,[row('SEI-1')]);wb=load_workbook(path);ws=wb['APÓLICES']
    table=Table(displayName='TeamTable',ref='A1:L2');ws.add_table(table);ws.auto_filter.ref='A1:L2'
    ws['M1']='MANUAL FORMULA';ws['M2']='=1+1';wb.save(path)
    update_workbook(path,item('SEI-1'),tmp_path/'backups')
    ws=load_workbook(path,data_only=False)['APÓLICES']
    assert ws.tables['TeamTable'].ref=='A1:L2' and ws.auto_filter.ref=='A1:L2' and ws['M2'].value=='=1+1'


def test_reordered_alias_headers_visual_changes_and_blank_rows_are_tolerated(tmp_path):
    from openpyxl.styles import Alignment
    path=tmp_path/'human.xlsx';wb=Workbook();ws=wb.active;ws.title='APÓLICES'
    headers=['RESPONSÁVEL',' Nº DO PROCESSO SEI ','NÚMERO DA CONCORRÊNCIA',' Órgão Contratante ','Empresa','Nº DO LOTE','Objeto','Início Vigência','Fim Vigência','Prêmio','Registro SUSEP','Linha digitável boleto','Observações']
    ws.append(headers)
    ws.append(['Ana','SEI-1','018/SEME/2026','Sub Parelheiros','THI Engenharia','01','manual 1',None,None,None,None,None,'não tocar'])
    ws.append([None]*len(headers))
    ws.append(['Bea','SEI-2','018/SEME/2026','Sub Parelheiros','THI Engenharia','02','manual 2',None,None,None,None,None,'preservar'])
    ws.column_dimensions['B'].width=43;ws.row_dimensions[4].height=39;ws['M4'].alignment=Alignment(wrap_text=True,vertical='top');ws.freeze_panes='C3';ws.auto_filter.ref='A1:M4';wb.save(path)
    update_workbook(path,item('SEI-2','02'),tmp_path/'backups')
    ws=load_workbook(path)['APÓLICES']
    assert ws.max_row==4 and ws['G3'].value is None and ws['A4'].value=='Bea' and ws['M4'].value=='preservar'
    assert ws.column_dimensions['B'].width==43 and ws.row_dimensions[4].height==39
    assert ws['M4'].alignment.wrap_text and ws.freeze_panes=='C3' and ws.auto_filter.ref=='A1:M4'

def test_missing_lot_column_is_added_once_and_two_lots_are_distinct(tmp_path):
    path=tmp_path/'no-lot.xlsx';headers=[h for h in ROBOT_MANAGED_COLUMNS if h!='LOTE']+['STATUS']
    wb=Workbook();ws=wb.active;ws.title='APÓLICES';ws.append(headers)
    # A pre-existing record without a lote remains untouched when another lote arrives.
    values=['Sub Parelheiros','THI Engenharia','018/SEME/2026','SEI-1',None,None,None,None,None,None,'PAGO'];ws.append(values);wb.save(path)
    update_workbook(path,item('SEI-1','01'),tmp_path/'backups')
    ws=load_workbook(path)['APÓLICES'];lot_col=next(c for c in range(1,ws.max_column+1) if str(ws.cell(1,c).value).strip().casefold()=='lote')
    assert ws.max_row==3 and ws.cell(2,lot_col).value is None and ws.cell(3,lot_col).value=='01'
    update_workbook(path,item('SEI-1','02'),tmp_path/'backups')
    ws=load_workbook(path)['APÓLICES'];assert ws.max_row==4 and [ws.cell(r,lot_col).value for r in (3,4)]==['01','02']


def test_missing_incoming_sei_does_not_fallback_to_row_with_known_sei(tmp_path):
    path=tmp_path/'known-sei.xlsx';workbook(path,[row('SEI-EXISTENTE')])
    result=update_workbook(path,item(None),tmp_path/'backups')
    ws=load_workbook(path)['APÓLICES']
    assert result=='ADICIONADO' and ws.max_row==3
    assert ws['D2'].value=='SEI-EXISTENTE' and ws['D3'].value is None
