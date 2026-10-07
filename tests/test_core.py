from datetime import date
from email.message import EmailMessage
from pathlib import Path
from pypdf import PdfWriter
from openpyxl import Workbook,load_workbook
import pytest
from app.file_manager import sanitize_filename,publish,build_destination,build_document_names
from app.models import PolicyData
from app.validator import validate_policy,require_minimum
from app.excel_apolices import update_workbook,ROBOT_MANAGED_COLUMNS
from app.excel_robot_control import RobotControl,HEADERS
from app.email_client import MailMessage
from app.main import process_message

def data(**kw):
    base=dict(empresa="THI Engenharia",empresa_normalizada="THI",tipo_empresa="THI",orgao="Órgão X",orgao_normalizado="ORGAO X",numero_concorrencia_original="01/2026",numero_concorrencia_normalizado="01-2026",vigencia_data_inicial=date(2026,1,2),vigencia_data_final=date(2027,1,1),confianca_geral=.95)
    base.update(kw);return PolicyData(**base)
def raw_result(**changes):
    fields={"empresa":{"valor":"THI Engenharia","fonte":"APOLICE","confianca":.98},"empresa_normalizada":{"valor":"THI","fonte":"APOLICE","confianca":.99},"orgao":{"valor":"Órgão X","fonte":"APOLICE","confianca":.95},"orgao_normalizado":"ORGAO X","numero_concorrencia_original":{"valor":"01/2026","fonte":"APOLICE","confianca":.95},"numero_concorrencia_normalizado":{"valor":"01-2026","fonte":"APOLICE","confianca":.95},"processo_sei":{"valor":None,"fonte":"NAO_IDENTIFICADO","confianca":.1},"lote":{"valor":None,"fonte":"NAO_IDENTIFICADO","confianca":.1},"par_coerente":True,"objeto":{"valor":"Seguro","fonte":"APOLICE","confianca":.9},"vigencia_data_inicial":{"valor":"2026-01-02","fonte":"APOLICE","confianca":.98},"vigencia_data_final":{"valor":"2027-01-01","fonte":"APOLICE","confianca":.98},"valor_premio":{"valor":1234.56,"fonte":"APOLICE","confianca":.95},"numero_registro_susep":{"valor":None,"fonte":"NAO_IDENTIFICADO","confianca":.2},"linha_digitavel_boleto":{"valor":"12345","fonte":"BOLETO","confianca":.95},"tipo_empresa":"THI","nome_pasta":None,"nome_apolice":None,"nome_boleto":None,"confianca_geral":.95,"campos_com_duvida":["processo_sei","numero_registro_susep"],"observacoes":None}
    fields.update(changes);return fields

def make_pdf(path:Path):
    writer=PdfWriter();writer.add_blank_page(width=200,height=200)
    with path.open('wb') as f:writer.write(f)
def test_sanitize_and_deterministic_names(tmp_path):
    d=data();assert sanitize_filename('a<b>:/\\c.pdf')=='a-b----c.pdf'
    a=build_destination(tmp_path,d);b=build_destination(tmp_path,d)
    assert a==b and build_document_names(d)==build_document_names(d)
def test_structured_evidence_and_validation():
    parsed=validate_policy(raw_result());require_minimum(parsed)
    assert parsed.empresa_normalizada=='THI' and parsed.evidencias['processo_sei'].fonte=='NAO_IDENTIFICADO'
    assert parsed.confidence_for('numero_concorrencia_normalizado')==.95
    with pytest.raises(Exception):validate_policy({**raw_result(),"vigencia_data_inicial":{"valor":"ontem","fonte":"APOLICE","confianca":.99}})
    with pytest.raises(Exception):validate_policy({**raw_result(),"vigencia_data_inicial":{"valor":"2027-01-02","fonte":"APOLICE","confianca":.99}})
    with pytest.raises(ValueError):
        bad=validate_policy({**raw_result(),"vigencia_data_inicial":{"valor":"2026-01-02","fonte":"APOLICE","confianca":.99},"vigencia_data_final":{"valor":"2026-01-01","fonte":"APOLICE","confianca":.99}})
        require_minimum(bad)
def test_null_evidence_and_low_confidence():
    result=raw_result();result['processo_sei']={"valor":None,"fonte":"NAO_IDENTIFICADO","confianca":.1}
    parsed=validate_policy(result);assert parsed.processo_sei is None
    result['vigencia_data_inicial']={"valor":"2026-01-02","fonte":"APOLICE","confianca":.2}
    with pytest.raises(ValueError):require_minimum(validate_policy(result))
    without_evidence=data()
    with pytest.raises(ValueError,match='Evidência documental ausente'):require_minimum(without_evidence)
def test_robot_control_status_attempts_retry_and_headers(tmp_path):
    ctl=RobotControl(tmp_path/'robot.xlsx',max_attempts=2)
    assert ctl.can_retry('m1','1')
    pid=ctl.begin(message_id='m1',uid='1',status='PROCESSANDO',modelo_ia='mock/model')
    assert pid and ctl.attempts('m1','1')==1 and ctl.find('m1')['status']=='PROCESSANDO'
    ctl.record(message_id='m1',uid='1',status='ERRO',erro='fail',pasta_destino='old-path')
    assert ctl.can_retry('m1','1')
    ctl.begin(message_id='m1',uid='1')
    assert ctl.find('m1')['erro'] is None and ctl.find('m1')['pasta_destino'] is None
    ctl.record(message_id='m1',uid='1',status='ERRO',erro='new failure')
    assert not ctl.can_retry('m1','1')
    ctl.record(message_id='m1',uid='1',status='SUCESSO',erro='stale')
    assert ctl.find('m1')['erro'] is None
    assert len(ctl.rows()[0])>=len(HEADERS)
def test_robot_control_duplicate_documents_require_pair_on_same_row(tmp_path):
    ctl=RobotControl(tmp_path/'r.xlsx');ctl.record(message_id='m1',uid='1',status='SUCESSO',hash_apolice='a',hash_boleto='b')
    assert ctl.already_processed('',{'a','b'})
    assert not ctl.already_processed('',{'a','other'})
def test_excel_manual_null_backup_and_explicit_sheet(tmp_path):
    path=tmp_path/'ops.xlsx';wb=Workbook();ws=wb.active;ws.title='APÓLICES';ws.append(ROBOT_MANAGED_COLUMNS+['STATUS','RESPONSÁVEL'])
    ws.append(['Órgão X','THI Engenharia','01/2026','SEI-1',None,None,None,None,None,None,None,'PAGO','BEA']);wb.save(path)
    update_workbook(path,data(processo_sei='SEI-1'),tmp_path/'backups')
    ws=load_workbook(path)['APÓLICES'];assert ws.cell(2,12).value=='PAGO' and ws.cell(2,13).value=='BEA'
    assert ws.cell(2,4).value=='SEI-1' and len(list((tmp_path/'backups').glob('*.xlsx')))==1
    update_workbook(path,data(processo_sei=None,objeto=None),tmp_path/'backups')
    ws=load_workbook(path)['APÓLICES'];assert ws.cell(2,4).value=='SEI-1' and ws.cell(2,12).value=='PAGO'
def test_excel_missing_named_sheet_and_ambiguous_rows_do_not_write(tmp_path):
    path=tmp_path/'bad.xlsx';wb=Workbook();wb.active.title='Other';wb.save(path)
    with pytest.raises(ValueError,match='APÓLICES'):update_workbook(path,data(),tmp_path/'backups')
    path2=tmp_path/'dup.xlsx';wb=Workbook();ws=wb.active;ws.title='APÓLICES';ws.append(ROBOT_MANAGED_COLUMNS)
    row=['Órgão X','THI Engenharia','01/2026',None,None,None,None,None,None,None];ws.append(row);ws.append(row);wb.save(path2)
    with pytest.raises(ValueError,match='POSSIVEL_DUPLICATA'):update_workbook(path2,data(),tmp_path/'backups')

def test_excel_different_sei_is_a_distinct_record_even_if_org_and_number_match(tmp_path):
    path=tmp_path/'sei.xlsx';wb=Workbook();ws=wb.active;ws.title='APÓLICES';ws.append(ROBOT_MANAGED_COLUMNS)
    ws.append(['Órgão X','THI Engenharia','01/2026','SEI-123',None,None,None,None,None,None]);wb.save(path)
    result=update_workbook(path,data(processo_sei='SEI-999'),tmp_path/'backups')
    ws=load_workbook(path)['APÓLICES']
    assert result=='ADICIONADO' and ws.max_row==3 and ws['D2'].value=='SEI-123' and ws['D3'].value=='SEI-999'
def test_excel_new_record_and_manual_columns(tmp_path):
    path=tmp_path/'new.xlsx';assert update_workbook(path,data(),tmp_path/'backups')=='ADICIONADO'
    ws=load_workbook(path)['APÓLICES'];assert ws.max_row==2 and ws['B2'].value=='THI Engenharia'
def test_publish_idempotent_and_conflict_rolls_back(tmp_path):
    p=tmp_path/'p.pdf';b=tmp_path/'b.pdf';make_pdf(p);make_pdf(b)
    dest,created=publish(tmp_path/'root',data(),p,b,return_created=True);assert len(created)==2
    dest2,created2=publish(tmp_path/'root',data(),p,b,return_created=True);assert dest==dest2 and not created2
    (dest/'08. BOLETO - ORGAO X - 01-2026.pdf').write_bytes(b'different')
    with pytest.raises(FileExistsError):publish(tmp_path/'root',data(),p,b)

def test_process_message_integration_mocked(tmp_path):
    attachments_dir=tmp_path/'source';attachments_dir.mkdir();policy=attachments_dir/'apolice.pdf';bill=attachments_dir/'boleto.pdf';make_pdf(policy);make_pdf(bill)
    em=EmailMessage();em['From']='Finlandia <operacao@finlandiaseguros.com.br>';em['To']='robot@example.com';em['Subject']='Documentos';em['Message-ID']='<m-1>';em.set_content('Segue documentação')
    em.add_attachment(policy.read_bytes(),maintype='application',subtype='pdf',filename='apolice.pdf');em.add_attachment(bill.read_bytes(),maintype='application',subtype='pdf',filename='boleto.pdf')
    msg=MailMessage('1','<m-1>','Finlandia <operacao@finlandiaseguros.com.br>','Documentos','2026-01-01',em.as_bytes())
    ops=tmp_path/'ops.xlsx';wb=Workbook();ws=wb.active;ws.title='APÓLICES';ws.append(ROBOT_MANAGED_COLUMNS+['STATUS']);ws.append(['Órgão X','THI Engenharia','01/2026',None,None,None,None,None,None,None,None,'PAGO']);wb.save(ops)
    class FakeClient:
        def classify(self,files,names):return {'grupos':[{'lote':{'valor':None,'fonte':'NAO_IDENTIFICADO','confianca':.1},'apolice':'apolice.pdf','boleto':'boleto.pdf'}],'outros':[],'observacoes':None}
        def analyze(self,policy_path,bill_path):return raw_result()
    ctl=RobotControl(tmp_path/'control.xlsx')
    status=process_message(msg,ctl,client=FakeClient(),root_dir=tmp_path/'docs',policies_excel=ops,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history')
    assert status=='SUCESSO';record=ctl.find('<m-1>');assert record['status']=='SUCESSO' and record['id_processamento']
    ws=load_workbook(ops)['APÓLICES'];assert ws['L2'].value=='PAGO' and ws['D2'].value is None
    assert (tmp_path/'history'/record['id_processamento']/'resultado.json').exists()
    assert len(list((tmp_path/'docs').rglob('*.pdf')))==2


def test_openrouter_sends_schema_and_parses_mocked_response(tmp_path,monkeypatch):
    import json
    from types import SimpleNamespace
    from app.openrouter_client import OpenRouterClient
    from app.prompt import POLICY_SCHEMA
    calls={}
    def fake_post(url,**kwargs):
        calls.update(kwargs)
        return SimpleNamespace(status_code=200,json=lambda:{"choices":[{"message":{"content":json.dumps(raw_result())}}]},raise_for_status=lambda:None)
    monkeypatch.setattr('app.openrouter_client.httpx.post',fake_post)
    client=OpenRouterClient('dummy-secret','mock/model','https://example.test/v1')
    f=tmp_path/'doc.pdf';make_pdf(f)
    response=client.analyze(f,f)
    assert response['empresa_normalizada']['valor']=='THI'
    schema=calls['json']['response_format']['json_schema']['schema']
    assert schema['properties']['processo_sei']['properties']['fonte']['enum']==['APOLICE','BOLETO','AMBOS','NAO_IDENTIFICADO']
    assert calls['headers']['HTTP-Referer']=='https://github.com/beasevla-hub/app_apolices'
    assert 'manus.im' not in str(calls['headers']).lower()
    assert 'dummy-secret' not in str(calls['json'])

def test_openrouter_invalid_json_retries_limited(tmp_path,monkeypatch):
    from types import SimpleNamespace
    from app.openrouter_client import OpenRouterClient,OpenRouterError
    count={'calls':0,'sleeps':0}
    def fake_post(*args,**kwargs):
        count['calls']+=1
        return SimpleNamespace(status_code=200,json=lambda:{"choices":[{"message":{"content":"not-json"}}]},raise_for_status=lambda:None)
    monkeypatch.setattr('app.openrouter_client.httpx.post',fake_post)
    monkeypatch.setattr('app.openrouter_client.time.sleep',lambda _:count.__setitem__('sleeps',count['sleeps']+1))
    client=OpenRouterClient('secret','mock/model','https://example.test/v1');f=tmp_path/'f.pdf';make_pdf(f)
    with pytest.raises(OpenRouterError):client.classify([f],[f.name])
    assert count=={'calls':3,'sleeps':2}

def test_phas_and_evidence_conflict_validation():
    result=raw_result();result['empresa_normalizada']={'valor':'PHAS','fonte':'APOLICE','confianca':.99};result['tipo_empresa']='PHAS'
    assert validate_policy(result).empresa_normalizada=='PHAS'
    result['tipo_empresa']='THI'
    with pytest.raises(ValueError):validate_policy(result)
