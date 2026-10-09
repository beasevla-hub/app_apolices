import json
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from pypdf import PdfWriter
from app.email_client import MailMessage,EmailClient
from app.excel_apolices import ROBOT_MANAGED_COLUMNS
from app.excel_robot_control import RobotControl
from app.main import process_message,CountingOpenRouter
from tests.test_core import raw_result

def pdf_bytes(width=200,height=200):
    writer=PdfWriter();writer.add_blank_page(width=width,height=height)
    import io
    buffer=io.BytesIO();writer.write(buffer);return buffer.getvalue()
def make_message(tmp_path,filenames,message_id='multi@example',payload=None,received_at='2026-09-25'):
    em=EmailMessage();em['From']='arquivo@finlandiaseguros.com.br';em['To']='robot@example.com';em['Subject']='Apólices de lotes';em['Message-ID']=f'<{message_id}>';em.set_content('Documentos para diversos lotes')
    payload=payload or pdf_bytes()
    for name in filenames:em.add_attachment(payload,maintype='application',subtype='pdf',filename=name)
    return MailMessage('445',f'<{message_id}>','arquivo@finlandiaseguros.com.br',em['Subject'],'2026-09-25',em.as_bytes(),received_at)
def setup_excel(path):
    wb=Workbook();ws=wb.active;ws.title='APÓLICES';ws.append(ROBOT_MANAGED_COLUMNS);wb.save(path)
class GroupClient:
    def __init__(self,groups,fail_once=None,coherence=None):self.groups=groups;self.fail_once=fail_once;self.failed=False;self.coherence=coherence or {};self.classifications=0;self.analyses=[]
    def classify(self,files,names):self.classifications+=1;return {'grupos':self.groups,'outros':['extra.pdf'] if 'extra.pdf' in names else [],'observacoes':None}
    def analyze(self,policy,bill,expected_lot=None):
        self.analyses.append((policy.name,bill.name,expected_lot))
        if expected_lot==self.fail_once and not self.failed:self.failed=True;raise RuntimeError('falha simulada do lote')
        coherent=self.coherence.get(expected_lot,True)
        evidence={'valor':expected_lot,'fonte':'APOLICE','confianca':.98} if expected_lot else {'valor':None,'fonte':'NAO_IDENTIFICADO','confianca':.1}
        return raw_result(lote=evidence,par_coerente=coherent)
def g(lot,policy,bill):return {'lote':{'valor':lot,'fonte':'APOLICE' if lot else 'NAO_IDENTIFICADO','confianca':.98 if lot else .1},'lotes':[{'valor':lot,'fonte':'APOLICE','confianca':.98}] if lot else [],'apolice':policy,'boleto':bill}
def run_process(tmp_path,message,client):
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx',max_attempts=4);counting=CountingOpenRouter(client)
    status=process_message(message,control,counting,root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history')
    return status,control,excel,counting

def test_two_lots_random_order_other_file_and_resume_only_failed_group_without_flag(tmp_path):
    names=['policy_02.pdf','bill_01.pdf','extra.pdf','policy_01.pdf','bill_02.pdf']
    message=make_message(tmp_path,names,'retry@example')
    client=GroupClient([g('01','policy_01.pdf','bill_01.pdf'),g('02','policy_02.pdf','bill_02.pdf')],fail_once='02')
    status,control,excel,counting=run_process(tmp_path,message,client)
    assert status=='ERRO' and counting.classifications==1 and counting.analyses==2
    rows=control.rows();success=[r for r in rows if (r.get('group_key') or '').startswith('PAIR-') and r.get('status')=='SUCESSO']
    assert len(success)==1
    # Retry reuses exact classification and skips the already-successful pair.
    status=process_message(message,control,counting,root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history')
    assert status=='SUCESSO' and counting.classifications==1 and counting.analyses==3
    ws=load_workbook(excel)['APÓLICES'];lot_col=next(c for c in range(1,ws.max_column+1) if ws.cell(1,c).value=='LOTE')
    assert {ws.cell(r,lot_col).value for r in range(2,ws.max_row+1)}=={'01','02'}
    folders=[p.name for p in (tmp_path/'docs'/'THI').rglob('LOTE *') if p.is_dir()]
    published=list((tmp_path/'docs').rglob('*.pdf'))
    assert set(folders)=={'LOTE 01','LOTE 02'} and len(published)==4
    assert {path.name for path in published}=={'01. APOLICE.pdf','08. BOLETO.pdf'}
    email_row=control.find(message.message_id,message.uid)
    metadata=json.loads((tmp_path/'history'/email_row['id_processamento']/'metadata.json').read_text())
    assert metadata['quantidade_pdfs']==5 and metadata['quantidade_grupos']==2
    assert len(metadata['pdfs'])==5 and all(len(item['sha256'])==64 for item in metadata['pdfs'])
    assert all(entry['status']=='SUCESSO' for entry in metadata['lotes'])
    assert metadata['classificacao_cache_reutilizada'] is True
    assert metadata['chamadas_openrouter']=={'classificacao':0,'analises':1,'total':1}
    # E-mail distinto classifica uma vez; hashes exatos dos pares concluídos evitam novas análises.
    copied=make_message(tmp_path,['policy_01.pdf','bill_01.pdf','policy_02.pdf','bill_02.pdf'],'copy@example')
    calls=(counting.classifications,counting.analyses)
    assert process_message(copied,control,counting,root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history')=='SUCESSO'
    assert (counting.classifications,counting.analyses)==(calls[0]+1,calls[1])
    assert process_message(message,control,counting,root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history')=='SUCESSO'
    assert (counting.classifications,counting.analyses)==(calls[0]+1,calls[1])

def test_ten_groups_make_ten_separate_analyses(tmp_path):
    names=[f'{kind}_{idx:02d}.pdf' for idx in range(1,11) for kind in ('apolice','boleto')]
    names=names[::2]+names[1::2][::-1]
    message=make_message(tmp_path,names,'ten@example')
    groups=[g(f'{idx:02d}',f'apolice_{idx:02d}.pdf',f'boleto_{idx:02d}.pdf') for idx in range(1,11)]
    status,control,excel,counting=run_process(tmp_path,message,GroupClient(groups))
    assert status=='SUCESSO' and counting.classifications==1 and counting.analyses==10
    assert load_workbook(excel)['APÓLICES'].max_row==11
    published=list((tmp_path/'docs').rglob('*.pdf'))
    assert len(published)==20 and {path.name for path in published}=={'01. APOLICE.pdf','08. BOLETO.pdf'}

def test_cross_lot_pair_rejected_by_independent_analysis(tmp_path):
    names=['policy_01.pdf','bill_01.pdf','policy_02.pdf','bill_02.pdf']
    message=make_message(tmp_path,names,'mismatch@example')
    groups=[g('01','policy_01.pdf','bill_02.pdf'),g('02','policy_02.pdf','bill_01.pdf')]
    client=GroupClient(groups,coherence={'01':False,'02':False})
    status,control,excel,counting=run_process(tmp_path,message,client)
    assert status=='ERRO' and counting.analyses==2
    assert not (tmp_path/'docs').exists() and load_workbook(excel)['APÓLICES'].max_row==1
    group_rows=[r for r in control.rows() if (r.get('group_key') or '').startswith('PAIR-')]
    assert len(group_rows)==2 and all(r['status']=='ERRO' for r in group_rows)


def test_other_only_message_is_ignored_without_policy_analysis(tmp_path):
    message=make_message(tmp_path,['extra-a.pdf','extra-b.pdf'],'other@example')
    class OtherClient:
        def __init__(self):self.classifications=0;self.analyses=0
        def classify(self,files,names):self.classifications+=1;return {'grupos':[],'outros':names,'observacoes':None}
        def analyze(self,*args,**kwargs):self.analyses+=1;raise AssertionError('não deve analisar OUTRO')
    status,control,excel,counting=run_process(tmp_path,message,OtherClient())
    assert status=='IGNORADO' and counting.classifications==1 and counting.analyses==0
    assert control.find(message.message_id,message.uid)['status']=='IGNORADO'
    assert not (tmp_path/'docs').exists() and load_workbook(excel)['APÓLICES'].max_row==1


def test_associated_lot_is_operational_when_policy_has_no_documentary_lot(tmp_path):
    message=make_message(tmp_path,['policy.pdf','bill.pdf'],'associated-only@example')
    class AssociatedOnlyClient:
        def classify(self,files,names):return {'grupos':[g('1','policy.pdf','bill.pdf')],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lot=None):
            return raw_result(lote={'valor':None,'fonte':'NAO_IDENTIFICADO','confianca':.1})
    status,control,excel,counting=run_process(tmp_path,message,AssociatedOnlyClient())
    assert status=='SUCESSO' and counting.classifications==1 and counting.analyses==1
    ws=load_workbook(excel)['APÓLICES'];lot_col=next(c for c in range(1,ws.max_column+1) if ws.cell(1,c).value=='LOTE')
    assert ws.cell(2,lot_col).value=='1'
    group_row=next(row for row in control.rows() if (row.get('group_key') or '').startswith('PAIR-'))
    assert group_row['lote']=='1' and group_row['status']=='SUCESSO'
    email=control.find(message.message_id,message.uid)
    metadata=json.loads((tmp_path/'history'/email['id_processamento']/'metadata.json').read_text())
    item=metadata['lotes'][0]
    assert item['lote']=='1' and item['lote_associado']=='1' and item['lote_documental'] is None
    assert len(list((tmp_path/'docs').rglob('LOTE *')))==0


def test_documentary_lot_conflict_with_associated_group_is_rejected(tmp_path):
    message=make_message(tmp_path,['policy.pdf','bill.pdf'],'lot-conflict@example')
    class ConflictingLotClient:
        def classify(self,files,names):return {'grupos':[g('1','policy.pdf','bill.pdf')],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lot=None):
            return raw_result(lote={'valor':'2','fonte':'APOLICE','confianca':.99})
    status,control,excel,counting=run_process(tmp_path,message,ConflictingLotClient())
    assert status=='ERRO' and counting.analyses==1
    row=next(row for row in control.rows() if (row.get('group_key') or '').startswith('PAIR-'))
    assert row['status']=='ERRO' and 'não confirma associação' in row['erro']
    assert not (tmp_path/'docs').exists() and load_workbook(excel)['APÓLICES'].max_row==1

@pytest.mark.parametrize('lots',[['01','02'],['01','02','03','04']])
def test_one_physical_pair_creates_one_row_per_lot_with_individual_premium_and_email_datetime(tmp_path,lots):
    names=['policy_multi.pdf','bill_multi.pdf'];message=make_message(tmp_path,names,'one-pair-many-lots@example',received_at='2026-10-09T11:31:42+00:00')
    evidence=lambda value:{'valor':value,'fonte':'APOLICE','confianca':.99}
    group={'lote':{'valor':None,'fonte':'NAO_IDENTIFICADO','confianca':.1},'lotes':[evidence(value) for value in lots],'apolice':names[0],'boleto':names[1]}
    class OnePairManyLotsClient:
        def __init__(self):self.classifications=0;self.analyses=[]
        def classify(self,files,filenames):
            self.classifications+=1
            return {'grupos':[group],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lots=None):
            self.analyses.append((policy.name,bill.name,tuple(expected_lots or ())))
            assert expected_lots==lots
            details=[{'numero':evidence(value),'valor_premio':{'valor':index*1000.25,'fonte':'APOLICE','confianca':.97}} for index,value in enumerate(lots,1)]
            return raw_result(lote={'valor':None,'fonte':'NAO_IDENTIFICADO','confianca':.1},lotes=details,valor_premio={'valor':None,'fonte':'NAO_IDENTIFICADO','confianca':.1})
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx');client=OnePairManyLotsClient();counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history')
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    ws=load_workbook(excel)['APÓLICES'];headers={ws.cell(1,c).value:c for c in range(1,ws.max_column+1)}
    assert ws.max_row==len(lots)+1
    assert [ws.cell(row,headers['LOTE']).value for row in range(2,ws.max_row+1)]==lots
    assert [ws.cell(row,headers['VALOR DO PRÊMIO']).value for row in range(2,ws.max_row+1)]==[index*1000.25 for index in range(1,len(lots)+1)]
    expected_received=datetime(2026,10,9,8,31,42)
    received=[ws.cell(row,headers['DATA DE RECEBIMENTO DO E-MAIL']).value for row in range(2,ws.max_row+1)]
    assert received==[expected_received]*len(lots)
    assert all(ws.cell(row,headers['DATA DE RECEBIMENTO DO E-MAIL']).number_format=='dd/mm/yyyy hh:mm:ss' for row in range(2,ws.max_row+1))
    published=list((tmp_path/'docs').rglob('*.pdf'))
    assert len(published)==2 and {path.name for path in published}=={'01. APOLICE.pdf','08. BOLETO.pdf'}
    assert client.classifications==1 and len(client.analyses)==1 and client.analyses[0][2]==tuple(lots)
    # Repetir a execução como BACKFILL preserva as linhas por lote e reutiliza o sucesso técnico do par.
    assert process_message(message,control,counting,mode='BACKFILL',**kwargs)=='SUCESSO'
    assert load_workbook(excel)['APÓLICES'].max_row==len(lots)+1
    assert client.classifications==1 and len(client.analyses)==1

def test_general_premium_is_not_repeated_on_each_multilot_row(tmp_path):
    lots=['01','02'];names=['policy_general.pdf','bill_general.pdf'];message=make_message(tmp_path,names,'general-premium@example')
    ev=lambda value,source='APOLICE',confidence=.98:{'valor':value,'fonte':source,'confianca':confidence}
    group={'lote':ev(None,'NAO_IDENTIFICADO',.1),'lotes':[ev(lot) for lot in lots],'apolice':names[0],'boleto':names[1]}
    class GeneralPremiumClient:
        def classify(self,files,filenames):return {'grupos':[group],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lots=None):
            return raw_result(lote=ev(None,'NAO_IDENTIFICADO',.1),lotes=[{'numero':ev(lot),'valor_premio':ev(None,'NAO_IDENTIFICADO',.1)} for lot in lots],valor_premio=ev(4000.0))
    excel=tmp_path/'ops.xlsx';client=GeneralPremiumClient()
    status,_,_,_=run_process(tmp_path,message,client)
    assert status=='SUCESSO'
    ws=load_workbook(excel)['APÓLICES'];premium_col=next(c for c in range(1,ws.max_column+1) if ws.cell(1,c).value=='VALOR DO PRÊMIO')
    assert ws.max_row==3 and [ws.cell(r,premium_col).value for r in (2,3)]==[None,None]

def test_analysis_expands_incomplete_single_lot_classification_and_updates_cache(tmp_path):
    lots=['01','02','03','04'];names=['policy_expand.pdf','bill_expand.pdf'];message=make_message(tmp_path,names,'expand-lots@example')
    ev=lambda value:{'valor':value,'fonte':'APOLICE','confianca':.99}
    class ExpandingClient:
        def __init__(self):self.classifications=0;self.analyses=0
        def classify(self,files,filenames):
            self.classifications+=1
            return {'grupos':[g('01',names[0],names[1])],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lot=None,expected_lots=None):
            self.analyses+=1
            return raw_result(lote=ev(None),lotes=[{'numero':ev(lot),'valor_premio':ev(index*250.0)} for index,lot in enumerate(lots,1)],valor_premio=ev(None))
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx');client=ExpandingClient();counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history')
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    ws=load_workbook(excel)['APÓLICES'];lot_col=next(c for c in range(1,ws.max_column+1) if ws.cell(1,c).value=='LOTE')
    assert [ws.cell(row,lot_col).value for row in range(2,6)]==lots
    cache=list((tmp_path/'history'/'classificacoes').glob('*.json'))[0]
    cached=json.loads(cache.read_text())
    assert cached['schema_version']>=2 and len(cached['resultado']['grupos'][0]['lotes'])==4
    assert cached['resultado']['grupos'][0]['lote']['valor'] is None
    assert client.classifications==1 and client.analyses==1
    assert process_message(message,control,counting,mode='BACKFILL',**kwargs)=='SUCESSO'
    assert load_workbook(excel)['APÓLICES'].max_row==5 and client.classifications==1 and client.analyses==1

def test_multilot_pair_is_not_published_when_analysis_omits_the_structured_lot_list(tmp_path):
    lots=['01','02'];names=['policy_no_lots.pdf','bill_no_lots.pdf'];message=make_message(tmp_path,names,'missing-analysis-lots@example')
    ev=lambda value,source='APOLICE',confidence=.98:{'valor':value,'fonte':source,'confianca':confidence}
    group={'lote':ev(None,'NAO_IDENTIFICADO',.1),'lotes':[ev(lot) for lot in lots],'apolice':names[0],'boleto':names[1]}
    class OmittingClient:
        def classify(self,files,filenames):return {'grupos':[group],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lots=None):return raw_result(lote=ev(None,'NAO_IDENTIFICADO',.1),lotes=[])
    status,control,excel,_=run_process(tmp_path,message,OmittingClient())
    assert status=='ERRO' and load_workbook(excel)['APÓLICES'].max_row==1
    assert not list((tmp_path/'docs').rglob('*.pdf'))
    failure=next(row for row in control.rows() if row.get('group_key'))
    assert failure['status']=='ERRO' and 'lista estruturada de lotes' in failure['erro']


def test_backfill_reuses_classification_and_reanalyzes_only_failed_group_without_flag(tmp_path):
    names=[f'{kind}_{idx:02d}.pdf' for idx in range(1,4) for kind in ('policy','bill')]
    message=make_message(tmp_path,names,'retry-errors-groups@example')
    groups=[g('01','policy_01.pdf','bill_01.pdf'),g('02','policy_02.pdf','bill_02.pdf'),g('03','policy_03.pdf','bill_03.pdf')]
    client=GroupClient(groups,fail_once='02')
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx',max_attempts=1);counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history',mode='BACKFILL')
    assert process_message(message,control,counting,**kwargs)=='ERRO'
    assert counting.classifications==1 and counting.analyses==3
    rows=[row for row in control.rows() if (row.get('group_key') or '').startswith('PAIR-')]
    assert sorted(row['status'] for row in rows)==['ERRO','SUCESSO','SUCESSO']
    attempts_before={row['group_key']:row['tentativas'] for row in rows}
    assert max(attempts_before.values())==1
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    # A classificação vem do cache e apenas o grupo 02 recebe uma nova análise.
    assert counting.classifications==1 and counting.analyses==4
    rows=[row for row in control.rows() if (row.get('group_key') or '').startswith('PAIR-')]
    assert all(row['status']=='SUCESSO' for row in rows)
    assert sorted(row['tentativas'] for row in rows)==[1,1,2]


def test_backfill_retries_recent_processing_and_hundred_attempts_by_group(tmp_path):
    from datetime import datetime
    names=[f'{kind}_{idx:02d}.pdf' for idx in range(1,4) for kind in ('policy','bill')]
    message=make_message(tmp_path,names,'stale-processing-groups@example')
    groups=[g('01','policy_01.pdf','bill_01.pdf'),g('02','policy_02.pdf','bill_02.pdf'),g('03','policy_03.pdf','bill_03.pdf')]
    client=GroupClient(groups,fail_once='02')
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx',max_attempts=1);counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history',mode='BACKFILL')
    assert process_message(message,control,counting,**kwargs)=='ERRO'
    rows=[row for row in control.rows() if (row.get('group_key') or '').startswith('PAIR-')]
    failed=next(row for row in rows if row['lote']=='02');successes={row['lote']:row for row in rows if row['status']=='SUCESSO'}
    parent=control.find(message.message_id,message.uid)
    old_group_process=failed['id_processamento'];old_parent_process=parent['id_processamento']
    recent_time=datetime.now().isoformat(timespec='seconds')
    control.record(message_id=message.message_id,uid=message.uid,status='PROCESSANDO',tentativas=100,data_ultima_tentativa=recent_time)
    control.record(message_id=message.message_id,uid=message.uid,group_key=failed['group_key'],status='PROCESSANDO',tentativas=100,data_ultima_tentativa=recent_time)
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    assert counting.classifications==1 and counting.analyses==4
    assert control.attempts(message.message_id,message.uid)==101
    assert control.attempts(message.message_id,message.uid,failed['group_key'])==101
    assert control.find(message.message_id,message.uid,failed['group_key'])['id_processamento']!=old_group_process
    assert control.find(message.message_id,message.uid)['id_processamento']!=old_parent_process
    assert (tmp_path/'history'/old_parent_process/'metadata.json').exists()
    old_metadata=json.loads((tmp_path/'history'/old_parent_process/'metadata.json').read_text())
    old_failed=next(item for item in old_metadata['lotes'] if item['group_key']==failed['group_key'])
    assert old_failed['status']=='ERRO'
    for lot,row in successes.items():
        current=control.find(message.message_id,message.uid,row['group_key'])
        assert current['status']=='SUCESSO' and current['tentativas']==1 and current['id_processamento']==row['id_processamento']


@pytest.mark.parametrize('mode',['NORMAL','BACKFILL'])
@pytest.mark.parametrize('status',['ERRO','PROCESSANDO','IGNORADO','ESTADO_ANTIGO'])
def test_execution_processes_every_non_success_state_even_after_100_attempts(tmp_path,status,mode):
    message=make_message(tmp_path,['policy.pdf','bill.pdf'],f'state-{status}@example')
    client=GroupClient([g('01','policy.pdf','bill.pdf')])
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx',max_attempts=1);counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history',mode=mode)
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    group_row=next(row for row in control.rows() if row.get('group_key'))
    control.record(message_id=message.message_id,uid=message.uid,status=status,tentativas=100)
    control.record(message_id=message.message_id,uid=message.uid,group_key=group_row['group_key'],status=status,tentativas=100)
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    assert counting.classifications==1 and counting.analyses==2
    final=control.find(message.message_id,message.uid,group_row['group_key'])
    assert final['status']=='SUCESSO' and final['tentativas']==101


def test_success_with_changed_hash_is_analyzed_again_and_file_conflict_is_safe(tmp_path):
    names=['policy.pdf','bill.pdf'];message=make_message(tmp_path,names,'changed-content@example')
    changed=make_message(tmp_path,names,'changed-content@example',payload=pdf_bytes(320,320))
    client=GroupClient([g('01','policy.pdf','bill.pdf')]);excel=tmp_path/'ops.xlsx';setup_excel(excel)
    control=RobotControl(tmp_path/'robot.xlsx');counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history',mode='BACKFILL')
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    assert process_message(changed,control,counting,**kwargs)=='ERRO'
    assert counting.classifications==2 and counting.analyses==2
    rows=[row for row in control.rows() if row.get('group_key')]
    assert sorted(row['status'] for row in rows)==['ERRO','SUCESSO']
    assert rows[0]['hash_apolice']!=rows[1]['hash_apolice']
    assert len(list((tmp_path/'docs').rglob('*.pdf')))==2


def test_keyboard_interrupt_is_recorded_and_next_backfill_resumes(tmp_path):
    message=make_message(tmp_path,['policy.pdf','bill.pdf'],'interrupt@example')
    client=GroupClient([g('01','policy.pdf','bill.pdf')])
    original=client.analyze
    def interrupt_once(*args,**kwargs):
        if not getattr(client,'interrupted',False):
            client.interrupted=True
            raise KeyboardInterrupt()
        return original(*args,**kwargs)
    client.analyze=interrupt_once
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx');counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history',mode='BACKFILL')
    with pytest.raises(KeyboardInterrupt):process_message(message,control,counting,**kwargs)
    assert control.find(message.message_id,message.uid)['status']=='ERRO'
    interrupted_group=next(row for row in control.rows() if row.get('group_key'))
    assert interrupted_group['status']=='ERRO' and 'Interrompido' in interrupted_group['erro']
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    assert counting.classifications==1 and counting.analyses==2
    assert control.attempts(message.message_id,message.uid,interrupted_group['group_key'])==2


def test_run_continues_backfill_after_one_email_fails(monkeypatch,tmp_path,capsys):
    from dataclasses import replace
    from datetime import date
    import app.main as main_module
    first=make_message(tmp_path,['policy_a.pdf','bill_a.pdf'],'first-fails@example')
    second=make_message(tmp_path,['policy_b.pdf','bill_b.pdf'],'second-succeeds@example')
    second.uid='446'
    class FakeMail:
        save_pdf_attachments=staticmethod(EmailClient.save_pdf_attachments)
        def __init__(self,*args,**kwargs):self.last_found_count=2;self.last_relevant_count=2
        def __enter__(self):return self
        def __exit__(self,*args):return False
        def messages_between(self,start,end):return iter([first,second])
    class FakeOpenRouter:
        def __init__(self):self.classifications=0;self.analyses=[]
        def classify(self,files,names):
            self.classifications+=1
            policy=next(name for name in names if name.startswith('policy_'))
            bill=next(name for name in names if name.startswith('bill_'))
            return {'grupos':[g('01',policy,bill)],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lot=None):
            self.analyses.append(policy.name)
            if policy.name.startswith('policy_a'):raise RuntimeError('falha isolada do primeiro e-mail')
            result=raw_result(lote={'valor':None,'fonte':'NAO_IDENTIFICADO','confianca':.1})
            suffix='A' if policy.name=='policy_a.pdf' else 'B'
            result['orgao']['valor']=f'Órgão {suffix}';result['orgao_normalizado']=f'ORGAO {suffix}'
            result['numero_concorrencia_original']['valor']=f'0{suffix}/2026'
            result['numero_concorrencia_normalizado']['valor']=f'0{suffix}-2026'
            return result
    fake_client=FakeOpenRouter();ops=tmp_path/'ops.xlsx';setup_excel(ops);control_path=tmp_path/'robot.xlsx'
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(main_module,'settings',replace(main_module.settings,email_user='user',email_password='pass',allowed_sender_domains=('@finlandiaseguros.com.br',),robot_control=control_path,root_dir=tmp_path/'docs',policies_excel=ops,dry_run=False,keep_success_temp=True))
    monkeypatch.setattr(main_module,'EmailClient',FakeMail);monkeypatch.setattr(main_module,'_client',lambda:fake_client)
    assert main_module.run(backfill_period=(date(2026,9,21),date(2026,10,7)))==0
    assert 'BACKFILL FINALIZADO' in capsys.readouterr().out
    assert fake_client.classifications==2 and fake_client.analyses==['policy_a_analise.pdf','policy_b_analise.pdf']
    control=RobotControl(control_path)
    assert control.find(first.message_id,first.uid)['status']=='ERRO'
    assert control.find(second.message_id,second.uid)['status']=='SUCESSO'
    assert len(list((tmp_path/'docs').rglob('*.pdf')))==2



def test_run_ctrl_c_stops_before_next_email_and_later_backfill_resumes(tmp_path,monkeypatch):
    from dataclasses import replace
    from datetime import date
    import app.main as main_module
    payload=pdf_bytes()
    first=make_message(tmp_path,['policy_a.pdf','bill_a.pdf'],'interrupt-run-a@example',payload=payload)
    second=make_message(tmp_path,['policy_b.pdf','bill_b.pdf'],'interrupt-run-b@example',payload=payload);second.uid='446'
    class FakeMail:
        save_pdf_attachments=staticmethod(EmailClient.save_pdf_attachments)
        def __init__(self,*args,**kwargs):self.last_found_count=2;self.last_relevant_count=2
        def __enter__(self):return self
        def __exit__(self,*args):return False
        def messages_between(self,start,end):return iter([first,second])
    class InterruptingOpenRouter:
        def __init__(self):self.classifications=0;self.analysis_calls=0;self.interrupt_next=True
        def classify(self,files,names):
            self.classifications+=1
            policy_name=next(name for name in names if name.startswith('policy_'))
            bill_name=next(name for name in names if name.startswith('bill_'))
            return {'grupos':[g('01',policy_name,bill_name)],'outros':[],'observacoes':None}
        def analyze(self,policy,bill,expected_lot=None):
            self.analysis_calls+=1
            if self.interrupt_next:
                self.interrupt_next=False
                raise KeyboardInterrupt()
            return raw_result(lote={'valor':'01','fonte':'APOLICE','confianca':.98})
    client=InterruptingOpenRouter();ops=tmp_path/'ops.xlsx';setup_excel(ops);control_path=tmp_path/'robot.xlsx'
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(main_module,'settings',replace(main_module.settings,email_user='user',email_password='pass',allowed_sender_domains=('@finlandiaseguros.com.br',),robot_control=control_path,root_dir=tmp_path/'docs',policies_excel=ops,dry_run=False,keep_success_temp=True))
    monkeypatch.setattr(main_module,'EmailClient',FakeMail);monkeypatch.setattr(main_module,'_client',lambda:client)
    period=(date(2026,9,21),date(2026,10,7))
    with pytest.raises(KeyboardInterrupt):main_module.run(backfill_period=period)
    control=RobotControl(control_path)
    assert client.classifications==1 and client.analysis_calls==1
    assert control.find(first.message_id,first.uid)['status']=='ERRO'
    interrupted_group=next(row for row in control.rows() if row.get('group_key'))
    assert interrupted_group['status']=='ERRO' and 'KeyboardInterrupt' in interrupted_group['erro']
    assert control.find(second.message_id,second.uid) is None
    first_history=tmp_path/'data'/'historico'/control.find(first.message_id,first.uid)['id_processamento']
    assert json.loads((first_history/'metadata.json').read_text())['status']=='ERRO'
    assert (first_history/'resultado.json').exists()

    assert main_module.run(backfill_period=period)==0
    assert client.analysis_calls==2  # a chamada interrompida não repetiu; só uma execução nova analisou o grupo.
    assert client.classifications==2  # a primeira mensagem reutiliza cache; classifica-se apenas a segunda.
    assert control.find(first.message_id,first.uid,interrupted_group['group_key'])['status']=='SUCESSO'
    assert control.find(second.message_id,second.uid)['status']=='SUCESSO'


def test_main_converts_keyboard_interrupt_to_short_exit_message(monkeypatch,capsys):
    import sys
    import app.main as main_module
    monkeypatch.setattr(sys,'argv',['app-apolices','--test'])
    monkeypatch.setattr(main_module,'_test',lambda:(_ for _ in ()).throw(KeyboardInterrupt()))
    assert main_module.main()==130
    captured=capsys.readouterr()
    assert 'Ctrl+C' in captured.err and 'Traceback' not in captured.err



def test_ctrl_c_in_later_group_preserves_earlier_success(tmp_path):
    message=make_message(tmp_path,['policy_01.pdf','bill_01.pdf','policy_02.pdf','bill_02.pdf'],'interrupt-second-group@example')
    class InterruptSecondGroup(GroupClient):
        def analyze(self,policy,bill,expected_lot=None):
            if expected_lot=='02' and not getattr(self,'interrupted',False):
                self.interrupted=True;self.analyses.append((policy.name,bill.name,expected_lot));raise KeyboardInterrupt()
            return super().analyze(policy,bill,expected_lot)
    client=InterruptSecondGroup([g('01','policy_01.pdf','bill_01.pdf'),g('02','policy_02.pdf','bill_02.pdf')])
    excel=tmp_path/'ops.xlsx';setup_excel(excel);control=RobotControl(tmp_path/'robot.xlsx');counting=CountingOpenRouter(client)
    kwargs=dict(root_dir=tmp_path/'docs',policies_excel=excel,backup_dir=tmp_path/'backups',temp_root=tmp_path/'temp',history_root=tmp_path/'history',mode='BACKFILL')
    with pytest.raises(KeyboardInterrupt):process_message(message,control,counting,**kwargs)
    group_rows={row['lote']:row for row in control.rows() if row.get('group_key')}
    first_success=group_rows['01'];second_interrupted=group_rows['02']
    assert first_success['status']=='SUCESSO' and first_success['tentativas']==1
    assert second_interrupted['status']=='ERRO' and 'KeyboardInterrupt' in second_interrupted['erro']
    assert process_message(message,control,counting,**kwargs)=='SUCESSO'
    after={row['lote']:row for row in control.rows() if row.get('group_key')}
    assert after['01']['status']=='SUCESSO' and after['01']['tentativas']==1
    assert after['01']['id_processamento']==first_success['id_processamento']
    assert after['02']['status']=='SUCESSO' and after['02']['tentativas']==2
    assert counting.classifications==1 and counting.analyses==3
